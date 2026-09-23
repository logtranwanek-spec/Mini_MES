"""One-time CNC migration. Dry run by default; --apply commits with SQLite backup.
Requires: py -m pip install openpyxl
Raw source rows and rejected records are retained; source workbook is never edited.
"""
import argparse, collections, datetime as dt, hashlib, json, pathlib, re, shutil, sqlite3
import openpyxl

def text(v):
    return '' if v is None else str(v).strip()

def date(v):
    if v is None or not text(v): return None
    if isinstance(v, dt.datetime): return v.strftime('%Y-%m-%d 00:00:00')
    for fmt in ('%d/%m/%Y', '%d/%m/%y', '%Y-%m-%d', '%m/%d/%Y'):
        try: return dt.datetime.strptime(text(v), fmt).strftime('%Y-%m-%d 00:00:00')
        except ValueError: pass
    raise ValueError('Ngày không hợp lệ: ' + text(v))

def time(v):
    if v is None or not text(v): return None
    if isinstance(v, dt.datetime): v = v.time()
    if isinstance(v, dt.time): return v.strftime('%H:%M:%S')
    if isinstance(v, (float, int)) and 0 <= v < 1:
        seconds = round(v * 86400)
        return f'{seconds//3600:02}:{seconds//60%60:02}:{seconds%60:02}'
    for fmt in ('%H:%M', '%H:%M:%S', '%H.%M'):
        try: return dt.datetime.strptime(text(v), fmt).strftime('%H:%M:%S')
        except ValueError: pass
    raise ValueError('Giờ không hợp lệ: ' + text(v))

def integer(v, name, low, high):
    try:
        n = float(v)
        if n.is_integer() and low <= n <= high: return int(n)
    except (ValueError, TypeError): pass
    raise ValueError(name + ' không hợp lệ: ' + text(v))

def convert(r, sheet):
    shift = text(r[0])
    if shift not in ('Day Shift', 'Night Shift'): raise ValueError('Ca không hợp lệ')
    machine = int(re.search(r'Heian\s*(\d+)', sheet, re.I)[1])
    cell_machine = re.search(r'\d+', text(r[3]))
    if cell_machine and int(cell_machine[0]) != machine: raise ValueError('Máy khác tên sheet')
    ins, rep = date(r[4]), date(r[9])
    if not ins: raise ValueError('Thiếu ngày lắp')
    if ins[:10] > dt.date.today().isoformat() or (rep and rep[:10] > dt.date.today().isoformat()):
        raise ValueError('Ngày trong tương lai')
    it, rt = time(r[5]), time(r[10])
    if rep and (rep < ins or (rep == ins and it and rt and rt < it)):
        raise ValueError('Thay trước khi lắp')
    pos = integer(r[6], 'Đầu dao', 1, 4)
    num = integer(r[7], 'Số dao', 1, 8)
    version = integer(r[8], 'Đợt cấp', 1, 100000)
    # A blank cell means the hour has not been entered yet; it is not zero hours.
    hours = integer(r[12], 'Giờ thực tế', 0, 24) if text(r[12]) else None
    if not rep and (hours is not None or text(r[13])): raise ValueError('Thiếu ngày thay nhưng có giờ/lý do thay')
    material, kind = text(r[14]), text(r[15])
    if 'mài' in sheet.lower(): material, kind = kind, material
    kind = kind.upper().replace('DAO MỚI', 'MỚI').replace('DAO MÀI', 'MÀI').replace('MÀI LÂN', 'MÀI LẦN')
    kind = re.sub(r'^MÀI (\d)$', r'MÀI LẦN \1', kind)
    reason = text(r[13])
    reason = {'cùn':'Cùn', 'gãy':'Gãy', 'mẻ':'Mẻ', 'cháy':'Cháy'}.get(reason.lower(), reason)
    return dict(Shift=shift, Supervisor=text(r[1]), MSS=text(r[2]), Date=ins,
        MachineName=f'Heian {machine}', ToolAddress=f'Heian{machine}-Dao{num}',
        ToolPosition=pos, ToolVersion=version, ToolType=kind, InstallDate=ins,
        InstallTime=it, ReplaceDate=rep, ReplaceTime=rt, ActualHours=hours,
        Reason=reason, Material=material.upper(), Supplier='', IsVersionIncrement=int(reason in ('Cùn','Gãy','Mẻ','Cháy')))

def main():
    p=argparse.ArgumentParser()
    p.add_argument('source'); p.add_argument('--db', default='Data/ToolManagement.db')
    p.add_argument('--apply', action='store_true'); a=p.parse_args()
    out=pathlib.Path('Data/CncImport'); out.mkdir(parents=True,exist_ok=True)
    source=pathlib.Path(a.source); digest=hashlib.sha256(source.read_bytes()).hexdigest()
    snapshot=out/(digest[:12]+'.xlsx')
    if not snapshot.exists(): shutil.copy2(source,snapshot)
    w=openpyxl.load_workbook(snapshot,read_only=True,data_only=True)
    raw=[]; valid=[]; errors=[]; counts=collections.Counter(); seen=set()
    for s in w:
        if not re.search(r'Heian\s*\d+',s.title,re.I): continue
        for n,r in enumerate(s.iter_rows(min_row=7,max_col=max(17,s.max_column),values_only=True),7):
            if not any(v is not None for v in r): continue
            rawjson=json.dumps(r,ensure_ascii=False,default=str)
            state=''; record=None
            if not any(text(r[i]) for i in (4,5,6,7,8,9,10,12)):
                state='template'; counts[state]+=1
            else:
                try:
                    record=convert(r,s.title)
                    key=json.dumps(record,sort_keys=True,ensure_ascii=False)
                    if key in seen: state='duplicate'
                    else: seen.add(key); state='valid'; valid.append((s.title,n,record))
                    counts[state]+=1
                except (ValueError,TypeError) as e:
                    state='review'; errors.append(dict(sheet=s.title,row=n,error=str(e),values=list(r)))
                    counts[state]+=1
            raw.append((digest,s.title,n,rawjson,state))
    w.close()
    report=dict(source=str(source),sha256=digest,counts=dict(counts),errors=errors)
    (out/'review.json').write_text(json.dumps(report,ensure_ascii=False,default=str,indent=2),encoding='utf-8')
    print(json.dumps(dict(counts=counts,error_types=collections.Counter(e['error'].split(':')[0] for e in errors)),ensure_ascii=False))
    if not a.apply: return
    c=sqlite3.connect(a.db,timeout=30); c.row_factory=sqlite3.Row
    backup=out/('ToolManagement.before-'+dt.datetime.now().strftime('%Y%m%d-%H%M%S')+'.db')
    with sqlite3.connect(backup) as b: c.backup(b)
    now=dt.datetime.now().isoformat(' ')
    inserted=0
    with c:
        c.execute('BEGIN IMMEDIATE')
        c.execute('CREATE TABLE IF NOT EXISTS CncExcelSourceRows (FileHash TEXT, Sheet TEXT, RowNumber INTEGER, RawJson TEXT, State TEXT, PRIMARY KEY(FileHash,Sheet,RowNumber))')
        c.execute('CREATE TABLE IF NOT EXISTS CncExcelImportedRecords (RecordHash TEXT PRIMARY KEY, ToolChangeId INTEGER, FileHash TEXT, Sheet TEXT, RowNumber INTEGER)')
        c.executemany('INSERT OR IGNORE INTO CncExcelSourceRows VALUES (?,?,?,?,?)',raw)
        fields=list(valid[0][2]) if valid else []
        existing={json.dumps({k:r[k] for k in fields},sort_keys=True,ensure_ascii=False) for r in c.execute('SELECT * FROM ToolChanges')}
        for sheet,n,r in valid:
            key=json.dumps(r,sort_keys=True,ensure_ascii=False); h=hashlib.sha256(key.encode()).hexdigest()
            if c.execute('SELECT 1 FROM CncExcelImportedRecords WHERE RecordHash=?',(h,)).fetchone() or key in existing: continue
            vals=dict(r,CreatedAt=now,UpdatedAt=now)
            cur=c.execute('INSERT INTO ToolChanges ('+','.join(vals)+') VALUES ('+','.join('?' for _ in vals)+')',list(vals.values()))
            c.execute('INSERT INTO CncExcelImportedRecords VALUES (?,?,?,?,?)',(h,cur.lastrowid,digest,sheet,n))
            existing.add(key); inserted+=1
        # Existing live statuses take precedence; initialize only missing addresses.
        groups=collections.defaultdict(list)
        for r in c.execute('SELECT * FROM ToolChanges'): groups[r['ToolAddress']].append(r)
        for address,rows in groups.items():
            latest=max(rows,key=lambda r:(r['InstallDate'] or r['Date'],r['InstallTime'] or '',r['Id']))
            c.execute('INSERT OR IGNORE INTO ToolStatuses (MachineName,ToolAddress,Shift,ToolPosition,CurrentVersion,TotalHours,CurrentVersionHours,LastUpdated) VALUES (?,?,?,?,?,?,?,?)',
                (latest['MachineName'],address,latest['Shift'],latest['ToolPosition'],latest['ToolVersion'],sum((r['ActualHours'] or 0) for r in rows),sum((r['ActualHours'] or 0) for r in rows if r['ToolVersion']==latest['ToolVersion']),now))
        assert c.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
    report.update(inserted=inserted,backup=str(backup),total_records=c.execute('SELECT count(*) FROM ToolChanges').fetchone()[0],statuses=c.execute('SELECT count(*) FROM ToolStatuses').fetchone()[0])
    (out/'result.json').write_text(json.dumps(report,ensure_ascii=False,default=str,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in report.items() if k!='errors'},ensure_ascii=False))
    c.close()

if __name__=='__main__': main()
